# tests/test_cross_lib.py

"""Cross-library parity: same markdown through `pdfmarq` and `docmarq` must both render."""

from pathlib import Path
import pytest
# every check here compares the two libraries, so without the twin there is nothing to run
pytest.importorskip("docmarq.legal", reason="docmarq 0.4 not installed")
from pdfmarq.paper import md_to_pdf, PaperStyle as PdfStyle
from pdfmarq.constants import A4 as PdfA4
from docmarq.paper import md_to_docx, PaperStyle as DocxStyle
from docmarq.constants import A4 as DocA4

#----------------------------------------------------------------------------------- Shared sources

_BASIC = "# Title\n\nFirst paragraph with **bold** and *italic*."
_TABLE = (
  "| A | B |\n|---|--:|\n| a | 1 |\n| b | 2 |\n"
)
_LISTS_CODE = (
  "## Recipe\n\n- prep\n- mix\n- bake\n\n"
  "```python\ndef hello():\n  return 42\n```\n"
)
_CALLOUT_QUOTE_HR = (
  "> [!NOTE]\n> Pay attention.\n\n> A normal quote.\n\n---\n\nAfter horizontal rule."
)
_FRONTMATTER = (
  "---\ntitle: Cross-lib doc\nauthor: Xaeian\n---\n\n"
  "# Body\n\nContent here."
)

#------------------------------------------------------------------------------------------ Helpers

@pytest.fixture
def is_pdf():
  def _check(path:Path) -> bool:
    with Path(path).open("rb") as f:
      return f.read(5) == b"%PDF-"
  return _check

@pytest.fixture
def is_docx():
  def _check(path:Path) -> bool:
    with Path(path).open("rb") as f:
      return f.read(4) == b"PK\x03\x04"
  return _check

@pytest.fixture
def render_both(tmp_path):
  def _render(src:str, name:str):
    pdf_path = tmp_path / f"{name}.pdf"
    docx_path = tmp_path / f"{name}.docx"
    md_to_pdf(src, str(pdf_path))
    md_to_docx(src, str(docx_path))
    return pdf_path, docx_path
  return _render

#------------------------------------------------------------------------------- Same-source render

@pytest.mark.parametrize("name, src", [
  ("basic", _BASIC),
  ("table", _TABLE),
  ("lists_code", _LISTS_CODE),
  ("callout_quote_hr", _CALLOUT_QUOTE_HR),
  ("frontmatter", _FRONTMATTER),
])
def same_source_renders_in_both_libs(render_both, is_pdf, is_docx, name, src):
  pdf_path, docx_path = render_both(src, name)
  assert is_pdf(pdf_path), f"PDF render failed for {name!r}"
  assert is_docx(docx_path), f"DOCX render failed for {name!r}"
  assert pdf_path.stat().st_size > 200
  assert docx_path.stat().st_size > 2000

def frontmatter_never_sets_geometry(tmp_path):
  # a leftover `render:` block is inert content in BOTH libs - geometry belongs to the caller.
  # guards the content/form split against a quiet regression
  src = "---\nrender:\n  landscape: true\n  page: A3\n---\n\n# Wide"
  pdf = md_to_pdf(src, str(tmp_path / "inert.pdf"))
  doc = md_to_docx(src, str(tmp_path / "inert.docx"))
  assert (pdf.page_width, pdf.page_height) == (PdfA4.width, PdfA4.height)
  assert (doc.page_width, doc.page_height) == (DocA4.width, DocA4.height)

def page_landscape_flips_both(tmp_path):
  # `page=A4.landscape()` is the only way to flip, and it works in both libs
  pdf = md_to_pdf("# Wide", str(tmp_path / "ls.pdf"), page=PdfA4.landscape())
  doc = md_to_docx("# Wide", str(tmp_path / "ls.docx"), page=DocA4.landscape())
  assert pdf.page_width > pdf.page_height, "pdfmarq did not flip"
  assert doc.page_width > doc.page_height, "docmarq did not flip"

def footnote_label_handled_in_both(tmp_path):
  # footnote renders without crashing default (HR) and labeled (H2) in both libs
  src = "Body[^1].\n\n[^1]: footnote text."
  md_to_pdf(src, str(tmp_path / "fn_default.pdf"))
  md_to_docx(src, str(tmp_path / "fn_default.docx"))
  md_to_pdf(src, str(tmp_path / "fn_labeled.pdf"), style=PdfStyle(footnote_label="References"))
  md_to_docx(src, str(tmp_path / "fn_labeled.docx"), style=DocxStyle(footnote_label="References"))
  for name in ("fn_default", "fn_labeled"):
    assert (tmp_path / f"{name}.pdf").exists()
    assert (tmp_path / f"{name}.docx").exists()

#------------------------------------------------------------------------------------ Visual parity

def smaller_size_function_parity():
  # both libs export `smaller_size` with identical ladder semantics
  from pdfmarq.utils import smaller_size as pdf_smaller
  from docmarq.utils import smaller_size as doc_smaller
  for body in (8, 9, 10, 11, 12, 14, 16, 18, 20, 22, 24):
    assert pdf_smaller(body) == doc_smaller(body), \
      f"body={body}: pdf={pdf_smaller(body)} doc={doc_smaller(body)}"

def heading_sizes_parity():
  # same h1..h6 sizes so markdown headings render at the same scale
  from docmarq.constants import Defaults as DocDef
  def sizes(s) -> tuple:
    return (s.h1_size, s.h2_size, s.h3_size, s.h4_size, s.h5_size, s.h6_size)
  assert sizes(PdfStyle()) == sizes(DocxStyle()) == tuple(DocDef.HEAD_SIZES)

def table_font_size_derivation_parity():
  # default markdown table cell size auto-derives from body via smaller_size
  from pdfmarq.utils import smaller_size
  from docmarq.styles import TableStyle
  p = PdfStyle()
  assert p.table_size is None # None triggers derivation
  assert smaller_size(p.body_size) == 10 # body=11 → 10
  assert TableStyle().font_size is None

def footnote_font_size_derivation_parity():
  # bibliography uses the same smaller_size(body) derivation in both libs
  from pdfmarq.utils import smaller_size as pdf_smaller
  from docmarq.utils import smaller_size as doc_smaller
  for body in (10, 11, 12, 14):
    assert pdf_smaller(body) == doc_smaller(body)

def markdown_body_line_height_intentionally_different():
  # INTENTIONAL: reportlab applies line_height as a direct ratio (pdf=1.4),
  # Word multiplies the font's natural height, so docmarq=1.0 lands comparable.
  # regression guard against well-meaning "consistency" cleanup
  assert PdfStyle().line_height == 1.4
  assert DocxStyle().line_height == 1.0

def mermaid_scale_parity():
  # same mermaid scale → same cache key + output regardless of pipeline
  assert PdfStyle().mermaid_scale == DocxStyle().mermaid_scale

def banner_title_size_parity():
  assert PdfStyle().banner_title_size == DocxStyle().banner_title_size

def callout_colors_parity():
  # same per-type RGB tuples for callout border + text in both libs
  assert PdfStyle().callout_colors == DocxStyle().callout_colors

def mark_bg_default_parity():
  # both libs default `mark_bg` to the same named highlight ('yellow')
  assert PdfStyle().mark_bg == DocxStyle().mark_bg == "yellow"

def block_image_box_parity():
  # both libs resolve a block image to the same box
  from pdfmarq.md.md_images import ImageInfo, size_block
  from docmarq.md.image_utils import ImageDSL, apply_dsl_dims
  for px_w, px_h, dpi in [(96, 48, 96), (1200, 800, 96), (1200, 800, 300),
      (3000, 2000, 600), (800, 2400, 96), (600, 400, 300)]:
    info = ImageInfo(src="x.png", is_svg=False, dpi=dpi,
      nat_w_mm=px_w * 25.4 / dpi, nat_h_mm=px_h * 25.4 / dpi)
    pdf_box = size_block(info, 170, 120)
    doc_box = apply_dsl_dims(px_w, px_h, 170, 120, ImageDSL(), dpi, 150)
    assert pdf_box == pytest.approx(doc_box, abs=0.05), \
      f"{px_w}x{px_h}@{dpi}: pdf={pdf_box} doc={doc_box}"

@pytest.mark.parametrize("count, rows",
  [(2, [2]), (7, [7]), (8, [4, 4]), (15, [5, 5, 5])])
def image_grid_parity(tmp_path, monkeypatch, count, rows):
  # an image-only paragraph lays out into the same grid, at the same widths
  from PIL import Image
  from docx import Document
  from pdfmarq import PDF
  for i in range(count):
    Image.new("RGB", (1200, 800), (100, 150, 200)).save(tmp_path / f"i{i}.png")
  md = "\n".join(f"![](i{i}.png)" for i in range(count))
  pdf_w = []
  orig = PDF.image
  def spy(self, src, w=None, h=None, *a, **k):
    pdf_w.append(w)
    return orig(self, src, w, h, *a, **k)
  monkeypatch.setattr(PDF, "image", spy)
  md_to_pdf(md, str(tmp_path / "grid.pdf"), base_dir=str(tmp_path))
  md_to_docx(md, str(tmp_path / "grid.docx"), base_dir=str(tmp_path))
  ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}ext"
  paras = [p._element.findall(f".//{ns}")
    for p in Document(str(tmp_path / "grid.docx")).paragraphs
    if p._element.findall(f".//{ns}")]
  assert [len(p) for p in paras] == rows
  docx_w = [int(e.get("cx")) / 36000 for row in paras for e in row]
  assert pdf_w == pytest.approx(docx_w, abs=0.05), f"pdf={pdf_w} docx={docx_w}"

def image_min_dpi_parity():
  assert PdfStyle().image_min_dpi == DocxStyle().image_min_dpi

def image_row_threshold_parity():
  # both libs drop back from a row to inline at the same slot width
  from pdfmarq.md.md_blocks import _ROW_MIN_SLOT_MM as pdf_min
  from docmarq.md.renderer import _ROW_MIN_SLOT_MM as doc_min
  assert pdf_min == doc_min

def lang_presets_do_not_set_footnote_label():
  # footnote label is OFF by default in all language presets in both libs
  from pdfmarq.paper import LANG_PRESETS as PdfPresets
  from docmarq.paper import LANG_PRESETS as DocPresets
  for lang, preset in PdfPresets.items():
    assert "footnote_label" not in preset, f"pdfmarq lang={lang!r} sets footnote_label"
  for lang, preset in DocPresets.items():
    assert "footnote_label" not in preset, f"docmarq lang={lang!r} sets footnote_label"

def font_body_intentionally_different():
  # INTENTIONAL: pdfmarq → Vera (DejaVu Sans when found), docmarq → Calibri (Word native).
  # regression guard against well-meaning "consistency" cleanup
  assert PdfStyle().font_body != DocxStyle().font_body
  assert PdfStyle().font_body == "Vera"
  assert DocxStyle().font_body == "Calibri"

def md_to_signatures_share_arg_names():
  # md_to_pdf and md_to_docx accept the same cross-lib keyword subset
  import inspect
  pdf_params = set(inspect.signature(md_to_pdf).parameters)
  doc_params = set(inspect.signature(md_to_docx).parameters)
  shared = {"md_text", "output_path", "style", "page", "margin", "gutter",
    "base_dir", "font_dir", "metadata"}
  assert not (shared - pdf_params), f"md_to_pdf missing: {shared - pdf_params}"
  assert not (shared - doc_params), f"md_to_docx missing: {shared - doc_params}"

@pytest.mark.parametrize("name, src", [
  ("absent", "# x"),
  ("empty block", "---\n---\n\n# B"),
  ("not a mapping", "---\n- a\n- b\n---\n\n# B"),
  ("broken yaml", "---\ntitle: [\n---\n\n# B"),
  ("unterminated", "---\ntitle: T\n\n# B"),
  ("crlf", "---\r\ntitle: T\r\n---\r\n\r\n# B"),
  ("values as written", "---\nid: 0012\nversion: 1.10\ndraft: false\nref:\n---\n\n# B"),
])
@pytest.mark.filterwarnings("ignore:invalid YAML frontmatter")
def frontmatter_readers_agree(name, src):
  # both packages must read the same document the same way, `None` included:
  # callers do `fm.get(...)`, and the document packages split it alike
  from pdfmarq.md import peek_frontmatter as pdf_peek, split_frontmatter as pdf_split
  from docmarq.md import peek_frontmatter as doc_peek, split_frontmatter as doc_split
  assert pdf_peek(src) == doc_peek(src), f"readers disagree on {name!r}"
  assert pdf_split(src) == doc_split(src), f"splits disagree on {name!r}"

def shared_helpers_are_identical():
  # helpers vendored twice on purpose; this is what keeps the copies from drifting
  import inspect
  import pdfmarq.paper.render as pdf_paper, docmarq.paper.render as doc_paper
  import pdfmarq.md.md_frontmatter as pdf_fm, docmarq.md.frontmatter as doc_fm
  import pdfmarq.parts as pdf_parts, docmarq.parts as doc_parts
  import pdfmarq.legal.header as pdf_header, docmarq.legal.header as doc_header
  import pdfmarq.legal.render as pdf_legal, docmarq.legal.render as doc_legal
  def same(a, b, name:str, *words:tuple[str, str]):
    left, right = inspect.getsource(getattr(a, name)), inspect.getsource(getattr(b, name))
    for pdf_word, docx_word in words:
      left = left.replace(pdf_word, docx_word)
    assert left == right, f"{name} drifted"
  same(pdf_paper, doc_paper, "_skip_matching_h1")
  same(pdf_fm, doc_fm, "metadata_from_frontmatter", ("PDF.metadata", "DOCX.metadata"))
  for name in ("default_sign_labels", "sign_labels"):
    same(pdf_parts, doc_parts, name)
  assert pdf_parts.LANG_LABELS == doc_parts.LANG_LABELS
  for name in ("lines_of", "_one_line"):
    same(pdf_header, doc_header, name)
  for name in ("_clauses", "_kind", "_preview", "_pair"):
    same(pdf_legal, doc_legal, name, ("md_html.is_", "is_"))

def legal_styles_match():
  # the same fields and defaults on both sides, so a legal script switches by its import
  import dataclasses
  from pdfmarq.legal import LegalStyle as PdfLegal, LANG_PRESETS as pdf_presets
  from docmarq.legal import LegalStyle as DocLegal, LANG_PRESETS as doc_presets
  from pdfmarq.md import MarkdownStyle as PdfEngine
  from docmarq.md import MarkdownStyle as DocEngine
  def own(style, engine) -> dict:
    base = {f.name for f in dataclasses.fields(engine)}
    return {f.name: getattr(style(), f.name) for f in dataclasses.fields(style)
      if f.name not in base}
  assert own(PdfLegal, PdfEngine) == own(DocLegal, DocEngine)
  for name in ("justify", "h1_size", "h2_size", "h1_underline", "h2_underline",
      "head_gap_top", "head_gap_bot", "para_gap", "keep_together"):
    assert getattr(PdfLegal(), name) == getattr(DocLegal(), name), name
  assert pdf_presets == doc_presets

def legal_entries_share_their_signatures():
  # page sizes are two classes alike, so defaults compare by their text
  import inspect
  from pdfmarq.legal import md_to_pdf as pdf_one, md_parallel_to_pdf as pdf_many
  from docmarq.legal import md_to_docx as doc_one, md_parallel_to_docx as doc_many
  def params(f) -> list[tuple]:
    return [(name, repr(p.default)) for name, p in inspect.signature(f).parameters.items()]
  assert params(pdf_one) == params(doc_one)
  assert params(pdf_many) == params(doc_many)

def column_solver_is_identical():
  # the same markdown table lays its columns out alike in both packages
  import inspect
  import pdfmarq.md.md_table as pdf_mod
  import docmarq.tables as doc_mod
  for name in ("_fit_columns", "_cap_widths"):
    pdf_src = inspect.getsource(getattr(pdf_mod, name))
    assert pdf_src == inspect.getsource(getattr(doc_mod, name)), f"{name} drifted"

def round_two_helpers_are_identical():
  # the reading, dating, checking and warning helpers both packages carry
  import inspect
  import pdfmarq.md.md_frontmatter as pdf_fm, docmarq.md.frontmatter as doc_fm
  import pdfmarq.utils as pdf_utils, docmarq.utils as doc_utils
  import pdfmarq._warn as pdf_warn, docmarq._warn as doc_warn
  import pdfmarq.legal.render as pdf_legal, docmarq.legal.render as doc_legal
  def same(a, b, name:str, *words:tuple[str, str]):
    left, right = inspect.getsource(getattr(a, name)), inspect.getsource(getattr(b, name))
    for pdf_word, docx_word in words:
      left = left.replace(pdf_word, docx_word)
    assert left == right, f"{name} drifted"
  same(pdf_fm, doc_fm, "_plain")
  same(pdf_fm, doc_fm, "read_yaml")
  same(pdf_fm, doc_fm, "_text_loader")
  same(pdf_utils, doc_utils, "format_date")
  same(pdf_warn, doc_warn, "caller_level", ("pdfmarq", "docmarq"))
  same(pdf_legal, doc_legal, "_checked", ("pdfmarq", "docmarq"),
    ("anything is drawn", "a word is written"))

#------------------------------------------------------------------------------------- Style fields

def engine_defaults_differ_only_where_documented():
  # each difference has a row in docmarq/md/readme.md; a new one needs a row there too
  import dataclasses
  from pdfmarq.md import MarkdownStyle as PdfEngine
  from docmarq.md import MarkdownStyle as DocEngine
  def names(cls) -> set[str]:
    return {f.name for f in dataclasses.fields(cls)}
  shared = names(PdfEngine) & names(DocEngine)
  differ = {n for n in shared if getattr(PdfEngine(), n) != getattr(DocEngine(), n)}
  assert differ == {"font_body", "font_head", "font_mono", "line_height", "keep_together",
    "math_fontset", "head_gap_top", "head_gap_bot", "hr_thick"}
  assert not names(DocEngine) - names(PdfEngine) # docmarq reads no field pdfmarq lacks

def paper_styles_match():
  # past the engine, the paper fields agree; pdfmarq alone draws the banner itself
  import dataclasses
  from pdfmarq.md import MarkdownStyle as PdfEngine
  from docmarq.md import MarkdownStyle as DocEngine
  from pdfmarq.paper import LANG_PRESETS as pdf_presets
  from docmarq.paper import LANG_PRESETS as doc_presets
  import pdfmarq.paper.style as pdf_style, docmarq.paper.style as doc_style
  def own(cls, engine) -> set[str]:
    return {f.name for f in dataclasses.fields(cls)} - {f.name for f in dataclasses.fields(engine)}
  pdf_own, doc_own = own(PdfStyle, PdfEngine), own(DocxStyle, DocEngine)
  assert not doc_own - pdf_own
  assert {n for n in pdf_own - doc_own if not n.startswith("banner_")} == set()
  for name in pdf_own & doc_own:
    assert getattr(PdfStyle(), name) == getattr(DocxStyle(), name), name
  assert pdf_presets == doc_presets
  assert pdf_style._default_status_colors() == doc_style._default_status_colors()
  assert pdf_style.UNKNOWN_STATUS == doc_style.UNKNOWN_STATUS

def paper_entries_share_their_signatures():
  # `md_into` takes the target first and docmarq a font root for diagrams, nothing else differs
  import inspect
  from pdfmarq.paper import md_into as pdf_into
  from docmarq.paper import md_into as doc_into
  def params(f) -> list[tuple]:
    return [(name, repr(p.default)) for name, p in inspect.signature(f).parameters.items()]
  assert params(md_to_pdf) == params(md_to_docx)
  assert params(pdf_into)[1:] == [p for p in params(doc_into)[1:] if p[0] != "font_dir"]

#-------------------------------------------------------------------------------------- Same output

def repeated_headings_number_alike():
  from markdown_it import MarkdownIt
  from pdfmarq.md.markdown import MarkdownRenderer as PdfRenderer
  from docmarq.md.slug import collect_heading_slugs
  src = "# Setup\n\n## Setup\n\n## `Setup` zażółć\n\n### Setup\n\n## Setup-1\n"
  tokens = MarkdownIt("commonmark").parse(src)
  assert PdfRenderer._collect_heading_slugs(tokens) == collect_heading_slugs(tokens)

def page_number_reads_alike(tmp_path):
  # pdfmarq prints the numbers, Word shows its fields: the same label, digits and slash
  fitz = pytest.importorskip("fitz")
  import re, zipfile
  md_to_pdf("Body.", str(tmp_path / "n.pdf"))
  md_to_docx("Body.", str(tmp_path / "n.docx"))
  page = fitz.open(str(tmp_path / "n.pdf"))[0].get_text()
  footer = zipfile.ZipFile(tmp_path / "n.docx").read("word/footer1.xml").decode("utf-8")
  docx_text = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", footer))
  assert "Page 1/1" in page
  assert re.sub(r"\d+", "N", docx_text) == "Page N/N"

#--------------------------------------------------------------------------------------- Mail merge

def merge_is_one_module_in_two_packages():
  # vendored twice on purpose: only the package name may differ
  pytest.importorskip("jinja2")
  import pdfmarq.merge, docmarq.merge
  pdf_dir, doc_dir = Path(pdfmarq.merge.__file__).parent, Path(docmarq.merge.__file__).parent
  for name in ("__init__.py", "template.py"):
    pdf_source = (pdf_dir / name).read_text(encoding="utf-8").replace("pdfmarq", "docmarq")
    assert pdf_source == (doc_dir / name).read_text(encoding="utf-8"), f"merge/{name} drifted"

def merge_fills_alike():
  pytest.importorskip("jinja2")
  from pdfmarq.merge import fill as pdf_fill
  from docmarq.merge import fill as doc_fill
  template = ('---\ntitle: "{{ t }}"\n---\n'
    "{% for v in vs %}\n- {{ v }} {{ k | number }}\n{% endfor %}\n")
  data = {"t": "Contract: A", "vs": ["(a)", "[b]", "1) c", ":d:"], "k": "1234,5"}
  assert pdf_fill(template, data, lang="pl") == doc_fill(template, data, lang="pl")
